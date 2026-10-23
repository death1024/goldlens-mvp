# -*- coding: utf-8 -*-
"""Shared primitives for the GoldLens 2.0 technical plan."""

import json
import os
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIAGRAMS = os.path.join(ROOT, 'docs', 'diagrams')
BODY_FONT = '微软雅黑'
SERIF_FONT = '宋体'
INK = RGBColor(0x1B, 0x1D, 0x18)
MUT = RGBColor(0x6F, 0x74, 0x68)
GOLD = RGBColor(0xB8, 0x89, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class Plan:
    def __init__(self):
        self.blocks = []
        self.sections = []

    def h1(self, num, title, subtitle=''):
        self.sections.append((num, title))
        self.blocks.append(('h1', num, title, subtitle))

    def h2(self, num, title):
        self.blocks.append(('h2', num, title))

    def p(self, text, lead=''):
        self.blocks.append(('p', text, lead))

    def bullets(self, items, numbered=False):
        self.blocks.append(('bullets', list(items), numbered))

    def table(self, caption, headers, rows, widths, font_size=8.2, align_center=None):
        self.blocks.append(('table', caption, headers, rows, widths, font_size, align_center or set()))

    def image(self, filename, caption, width=16.3):
        self.blocks.append(('image', filename, caption, width))

    def callout(self, label, text):
        self.blocks.append(('callout', label, text))

    def code(self, text):
        self.blocks.append(('code', text))

    def quote(self, text):
        self.blocks.append(('quote', text))

    def pagebreak(self):
        self.blocks.append(('pagebreak',))


# ---------------- DOCX helpers ----------------

def set_font(run, name=BODY_FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(paragraph, text, size, color=INK):
    parts = str(text).split('**')
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_font(run, size=size, color=color, bold=(idx % 2 == 1))


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def cell_margins(cell, top=65, bottom=65, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for tag, value in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = tcMar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def cell_borders(cell, color='DED8C8', size=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ['top', 'left', 'bottom', 'right']:
        tag = 'w:{}'.format(edge)
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), str(size))
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(end)


def add_header_footer(section, header_text='GOLDLENS 2.0  /  TECHNICAL PLAN  /  2026.10'):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ''
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run(header_text)
    set_font(r, size=7.2, bold=True, color=GOLD)

    fp = section.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('Proposed · for discussion     ')
    set_font(r, size=7, color=MUT)
    add_page_number(fp)


def page_number_start(section, start=1):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), str(start))


def render_docx(plan, output, page_map=None, cover='tech-cover.png', header_text='GOLDLENS 2.0  /  TECHNICAL PLAN  /  2026.10', toc_title='技术方案 · 目录', doc_title='GoldLens 2.0 技术方案'):
    page_map = page_map or {}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for margin in ['left_margin', 'right_margin', 'top_margin', 'bottom_margin']:
        setattr(sec, margin, Cm(0))
    sec.header_distance = Cm(0)
    sec.footer_distance = Cm(0)

    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), BODY_FONT)
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)

    # cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_picture(os.path.join(DIAGRAMS, cover), width=Cm(21.0), height=Cm(29.65))

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.page_width, body.page_height = Cm(21), Cm(29.7)
    body.left_margin = body.right_margin = Cm(2.1)
    body.top_margin = Cm(2.15)
    body.bottom_margin = Cm(1.9)
    body.header_distance = Cm(0.8)
    body.footer_distance = Cm(0.8)
    add_header_footer(body, header_text)
    page_number_start(body, 1)

    # TOC
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run('DOCUMENT MAP')
    set_font(r, size=8.5, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(toc_title)
    set_font(r, name=SERIF_FONT, size=27, bold=True, color=INK)
    for num, title in plan.sections:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.5), WD_TAB_ALIGNMENT.LEFT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(num)
        set_font(r, name='Georgia', size=8.5, bold=True, color=GOLD)
        r = p.add_run('\t' + title)
        set_font(r, size=10.3, color=INK)
        page = page_map.get(title)
        r = p.add_run('\t' + (str(page) if page else '—'))
        set_font(r, name='Georgia', size=9, color=MUT)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('阅读说明')
    set_font(r, size=8, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run('本方案面向创始人、工程与设计团队，描述从当前 MVP 到生产级 GoldLens 的目标架构、数据模型、体验规范与交付路线。所有时间与成本为显式假设，不是承诺；上线前需完成数据授权、隐私与本地合规审查。')
    set_font(r, size=8.8, color=MUT)

    doc.add_page_break()

    for block in plan.blocks:
        kind = block[0]
        if kind == 'pagebreak':
            doc.add_page_break()
            continue
        if kind == 'h1':
            _, num, title, subtitle = block
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(num)
            set_font(r, name='Georgia', size=9.5, bold=True, color=GOLD)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(title)
            set_font(r, name=SERIF_FONT, size=23, bold=True, color=INK)
            p.paragraph_format.keep_with_next = True
            if subtitle:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(10)
                r = p.add_run(subtitle)
                set_font(r, size=9.5, color=MUT)
                p.paragraph_format.keep_with_next = True
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '8')
            bottom.set(qn('w:color'), 'B8892E')
            pBdr.append(bottom)
            pPr.append(pBdr)
            continue
        if kind == 'h2':
            _, num, title = block
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(num + '  ')
            set_font(r, name='Georgia', size=9.5, bold=True, color=GOLD)
            r = p.add_run(title)
            set_font(r, size=12.3, bold=True, color=INK)
            p.paragraph_format.keep_with_next = True
            continue
        if kind == 'p':
            _, text, lead = block
            text = str(text).replace('`', '')
            lead = str(lead).replace('`', '')
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.42
            if lead:
                r = p.add_run(lead)
                set_font(r, size=10.2, bold=True, color=INK)
            add_inline(p, text, 10.2, INK)
            continue
        if kind == 'bullets':
            _, items, numbered = block
            for idx, item in enumerate(items):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.45
                if numbered:
                    r = p.add_run(f'{idx + 1}.  ')
                    set_font(r, name='Georgia', size=9.5, bold=True, color=GOLD)
                else:
                    r = p.add_run('·  ')
                    set_font(r, size=9.5, bold=True, color=GOLD)
                add_inline(p, str(item).replace('`', ''), 9.9, INK)
            continue
        if kind == 'table':
            _, caption, headers, rows, widths, font_size, align_center = block
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(caption)
            set_font(r, size=7.6, color=MUT)
            p.paragraph_format.keep_with_next = True
            table = doc.add_table(rows=1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for i, width in enumerate(widths):
                table.columns[i].width = Cm(width)
            for i, header in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.width = Cm(widths[i])
                shade(cell, '1B1D18')
                cell_margins(cell, 42, 42, 70, 70)
                cell_borders(cell, '1B1D18', 4)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in align_center else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(header)
                set_font(r, size=font_size, bold=True, color=WHITE)
            for ridx, row in enumerate(rows):
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cell = cells[i]
                    cell.width = Cm(widths[i])
                    cell_margins(cell, 30, 30, 70, 70)
                    cell_borders(cell, 'E2DCCD', 3)
                    if ridx % 2 == 1:
                        shade(cell, 'F8F5EC')
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.2
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in align_center else WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(str(value))
                    set_font(r, size=font_size, color=INK)
            for row in table.rows:
                trPr = row._tr.get_or_add_trPr()
                trPr.append(OxmlElement('w:cantSplit'))
            trPr = table.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:tblHeader'))
            if len(rows) <= 8:
                for row in table.rows[:-1]:
                    trPr = row._tr.get_or_add_trPr()
                    trPr.append(OxmlElement('w:keepNext'))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if kind == 'image':
            _, filename, caption, width = block
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(os.path.join(DIAGRAMS, filename), width=Cm(width))
            cp = doc.add_paragraph()
            cp.paragraph_format.space_after = Pt(7)
            r = cp.add_run(caption)
            set_font(r, size=7.6, color=MUT)
            continue
        if kind == 'callout':
            _, label, text = block
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            cell = table.cell(0, 0)
            cell.width = Cm(16.6)
            shade(cell, 'F3EBD8')
            cell_borders(cell, 'D9C48F', 6)
            cell_margins(cell, 120, 120, 160, 160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(label)
            set_font(r, size=7.8, bold=True, color=GOLD)
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.4
            r = p.add_run(text)
            set_font(r, size=9.5, color=INK)
            trPr = table.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if kind == 'code':
            _, text = block
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            cell = table.cell(0, 0)
            cell.width = Cm(16.6)
            shade(cell, 'F4F1E8')
            cell_borders(cell, 'DED8C8', 3)
            cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(text)
            set_font(r, name='Consolas', size=8.1, color=INK)
            trPr = table.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if kind == 'quote':
            _, text = block
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.left_indent = Cm(0.4)
            r = p.add_run(text)
            set_font(r, name=SERIF_FONT, size=18, bold=True, color=GOLD)
            continue

    doc.core_properties.title = doc_title
    doc.core_properties.author = 'GoldLens Research'
    doc.core_properties.subject = '从 MVP 到生产级黄金决策平台的技术方案'
    doc.core_properties.keywords = 'GoldLens, architecture, Next.js, Postgres, technical plan'
    doc.save(output)
    return output


def _md_escape(text):
    return str(text).replace('|', '\\|').replace('\n', ' ')


def render_markdown(plan, output, title='GoldLens 2.0 技术方案', subtitle='从 MVP 到可用的现代黄金决策平台。', version='Version 1.0 · Proposed · 2026-10-23'):
    lines = [
        f'# {title}',
        '',
        f'> {subtitle}',
        '',
        f'**{version}**',
        '',
        '---',
        '',
        '## 文档地图',
        '',
    ]
    for num, title in plan.sections:
        lines.append(f'- **{num}** {title}')
    lines.extend(['', '---', ''])

    for block in plan.blocks:
        kind = block[0]
        if kind == 'pagebreak':
            lines.extend(['', '---', ''])
        elif kind == 'h1':
            _, num, title, subtitle = block
            lines.append(f'## {num} {title}' if num != '00' else f'## {title}')
            if subtitle:
                lines.append(f'*{subtitle}*')
            lines.append('')
        elif kind == 'h2':
            _, num, title = block
            lines.append(f'### {num} {title}')
            lines.append('')
        elif kind == 'p':
            _, text, lead = block
            lines.append(f'**{lead}**{text}' if lead else text)
            lines.append('')
        elif kind == 'bullets':
            _, items, numbered = block
            for idx, item in enumerate(items):
                lines.append(f'{idx + 1}. {item}' if numbered else f'- {item}')
            lines.append('')
        elif kind == 'table':
            _, caption, headers, rows, widths, font_size, align_center = block
            lines.append(f'**{caption}**')
            lines.append('')
            lines.append('| ' + ' | '.join(_md_escape(h) for h in headers) + ' |')
            lines.append('|' + '|'.join(['---'] * len(headers)) + '|')
            for row in rows:
                lines.append('| ' + ' | '.join(_md_escape(v) for v in row) + ' |')
            lines.append('')
        elif kind == 'image':
            _, filename, caption, width = block
            lines.append(f'![{caption}](diagrams/{filename})')
            lines.append('')
            lines.append(f'*{caption}*')
            lines.append('')
        elif kind == 'callout':
            _, label, text = block
            lines.append(f'> **{label}**  ')
            lines.append(f'> {text}')
            lines.append('')
        elif kind == 'code':
            _, text = block
            lines.append('```text')
            lines.append(text)
            lines.append('```')
            lines.append('')
        elif kind == 'quote':
            _, text = block
            lines.append(f'> **{text}**')
            lines.append('')

    with open(output, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(lines))
    return output
